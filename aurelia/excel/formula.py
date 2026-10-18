"""A documented, deliberately limited Excel formula calculator for grading workbooks.

Why: openpyxl reads formulas as text and never calculates them, and a workbook
saved by a program other than Excel may contain no cached values (or stale ones).
Aurelia therefore recalculates supported formulas itself.

Supported
---------
* Literals: numbers, "text", TRUE/FALSE, percentages (10%).
* References: A1, $A$1, A1:B10, A:A (whole column, bounded by the sheet's used range),
  Sheet!A1, 'Sheet name'!A1:B2.
* Operators: + − * / ^ & = <> < > <= >= and unary minus.
* Functions: SUM, SUMIF, SUMIFS, COUNT, COUNTA, COUNTIF, COUNTIFS, AVERAGE, AVERAGEIF,
  AVERAGEIFS, MIN, MAX, ABS, ROUND, ROUNDUP, ROUNDDOWN, IF, IFERROR, AND, OR, NOT,
  XLOOKUP, VLOOKUP, INDEX, MATCH, SUMPRODUCT, MONTH, YEAR, DAY, DATE, EOMONTH, NPV, IRR, PMT.

Anything else raises :class:`UnsupportedFormula`; the grader then reports that the cell
could not be recalculated rather than guessing. This is NOT a full Excel engine and the
application does not claim Excel compatibility beyond this list.
"""
from __future__ import annotations

import calendar
import datetime as dt
import math
import re
from dataclasses import dataclass
from typing import Any

from openpyxl.utils import column_index_from_string, get_column_letter

SUPPORTED_FUNCTIONS = {"SUM", "SUMIF", "SUMIFS", "COUNT", "COUNTA", "COUNTIF", "COUNTIFS", "AVERAGE", "AVERAGEIF",
                       "AVERAGEIFS", "MIN", "MAX", "ABS", "ROUND", "ROUNDUP", "ROUNDDOWN", "IF", "IFERROR", "AND",
                       "OR", "NOT", "XLOOKUP", "VLOOKUP", "INDEX", "MATCH", "SUMPRODUCT", "MONTH", "YEAR", "DAY",
                       "DATE", "EOMONTH", "NPV", "IRR", "PMT"}


class UnsupportedFormula(Exception):
    pass


class FormulaSyntaxError(Exception):
    pass


@dataclass(frozen=True)
class XLError:
    code: str

    def __str__(self) -> str:
        return self.code


DIV0, NA, VALUE, REF, NAME, NUM = (XLError(c) for c in ("#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "#NUM!"))


class _Raise(Exception):
    def __init__(self, err: XLError):
        self.err = err


@dataclass
class Range:
    rows: list[list[Any]]

    def flat(self) -> list[Any]:
        return [v for r in self.rows for v in r]

    @property
    def shape(self):
        return len(self.rows), (len(self.rows[0]) if self.rows else 0)


# --------------------------------------------------------------------------- tokenizer
_TOKEN = re.compile(r"""
 (?P<ws>\s+)
|(?P<string>"(?:[^"]|"")*")
|(?P<error>\#(?:DIV/0!|N/A|VALUE!|REF!|NAME\?|NUM!|NULL!))
|(?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_\.]*)!)?
    (?:\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}))
|(?P<number>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?|\.\d+(?:[eE][+-]?\d+)?)
|(?P<func>[A-Za-z_][A-Za-z0-9_\.]*(?=\())
|(?P<bool>TRUE|FALSE|true|false)
|(?P<op><>|<=|>=|[-+*/^&=<>%])
|(?P<lp>\()|(?P<rp>\))|(?P<comma>[,;])
""", re.VERBOSE)


def tokenize(src: str) -> list[tuple[str, str]]:
    out, pos = [], 0
    while pos < len(src):
        m = _TOKEN.match(src, pos)
        if not m:
            raise FormulaSyntaxError(f"Unexpected character at position {pos}: {src[pos:pos + 10]!r}")
        kind = m.lastgroup
        text = m.group(kind)
        pos = m.end()
        if kind == "ws":
            continue
        out.append((kind, text))
    return out


# --------------------------------------------------------------------------- parser (Pratt)
_BINARY = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}


class Parser:
    def __init__(self, tokens):
        self.t = tokens
        self.i = 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self):
        tok = self.peek()
        self.i += 1
        return tok

    def parse(self):
        node = self.expr(0)
        if self.i != len(self.t):
            raise FormulaSyntaxError(f"Unexpected token {self.peek()[1]!r}")
        return node

    def expr(self, min_prec):
        left = self.unary()
        while True:
            kind, text = self.peek()
            if kind == "op" and text == "%":
                self.take()
                left = ("pct", left)
                continue
            if kind != "op" or text not in _BINARY or _BINARY[text] < min_prec:
                break
            prec = _BINARY[text]
            self.take()
            right = self.expr(prec + 1)  # Excel: all binary operators left-associative
            left = ("bin", text, left, right)
        return left

    def unary(self):
        kind, text = self.peek()
        if kind == "op" and text in "+-":
            self.take()
            return ("neg", self.unary()) if text == "-" else self.unary()
        return self.primary()

    def primary(self):
        kind, text = self.take()
        if kind == "number":
            return ("num", float(text))
        if kind == "string":
            return ("str", text[1:-1].replace('""', '"'))
        if kind == "bool":
            return ("bool", text.upper() == "TRUE")
        if kind == "error":
            return ("err", XLError(text))
        if kind == "ref":
            return ("ref", text)
        if kind == "func":
            name = text.upper()
            if name.startswith("_XLFN."):
                name = name[6:]
            self.take()  # (
            args = []
            if self.peek()[0] != "rp":
                while True:
                    if self.peek()[0] in ("comma", "rp"):
                        args.append(("missing",))
                    else:
                        args.append(self.expr(0))
                    if self.peek()[0] == "comma":
                        self.take()
                        continue
                    break
            if self.take()[0] != "rp":
                raise FormulaSyntaxError(f"Missing ')' after {name}")
            return ("func", name, args)
        if kind == "lp":
            node = self.expr(0)
            if self.take()[0] != "rp":
                raise FormulaSyntaxError("Missing ')'")
            return node
        raise FormulaSyntaxError(f"Unexpected token {text!r}")


def parse(formula: str):
    src = formula[1:] if formula.startswith("=") else formula
    return Parser(tokenize(src)).parse()


def functions_used(formula: str) -> set[str]:
    try:
        toks = tokenize(formula[1:] if formula.startswith("=") else formula)
    except FormulaSyntaxError:
        return set()
    out = set()
    for k, t in toks:
        if k == "func":
            n = t.upper()
            out.add(n[6:] if n.startswith("_XLFN.") else n)
    return out


def references_used(formula: str) -> list[str]:
    try:
        toks = tokenize(formula[1:] if formula.startswith("=") else formula)
    except FormulaSyntaxError:
        return []
    return [t for k, t in toks if k == "ref"]


# --------------------------------------------------------------------------- helpers
def _num(v):
    if isinstance(v, XLError):
        raise _Raise(v)
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, dt.datetime) or isinstance(v, dt.date):
        return float(_to_serial(v))
    if isinstance(v, str):
        try:
            return float(v.replace(",", ""))
        except ValueError:
            raise _Raise(VALUE)
    raise _Raise(VALUE)


_EPOCH = dt.datetime(1899, 12, 30)


def _to_serial(d) -> float:
    if isinstance(d, dt.date) and not isinstance(d, dt.datetime):
        d = dt.datetime(d.year, d.month, d.day)
    return (d - _EPOCH).days + (d - _EPOCH).seconds / 86400


def _to_date(v) -> dt.datetime:
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    return _EPOCH + dt.timedelta(days=_num(v))


def _cmp_key(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return _to_serial(v)
    return v


def _match_criterion(value, criterion) -> bool:
    """Excel criteria: 5, ">5", "<>x", "=x", "text", wildcards * and ?."""
    if isinstance(criterion, (int, float)) and not isinstance(criterion, bool):
        try:
            return _num(value) == float(criterion) if value not in (None, "") else False
        except _Raise:
            return False
    if isinstance(criterion, (dt.datetime, dt.date)):
        return value is not None and _cmp_key(value) == _cmp_key(criterion)
    c = "" if criterion is None else str(criterion)
    op = "="
    for o in ("<>", ">=", "<=", "=", ">", "<"):
        if c.startswith(o):
            op, c = o, c[len(o):]
            break
    try:
        cn = float(c)
        is_num = True
    except ValueError:
        is_num = False
    if is_num and value not in (None, "") and not isinstance(value, str):
        v = _cmp_key(value)
        return {"=": v == cn, "<>": v != cn, ">": v > cn, "<": v < cn, ">=": v >= cn, "<=": v <= cn}[op]
    sv = "" if value is None else str(value)
    if op in ("=", "<>"):
        pat = "^" + re.escape(c.lower()).replace(r"\*", ".*").replace(r"\?", ".") + "$"
        hit = re.match(pat, sv.lower()) is not None
        return hit if op == "=" else not hit
    if is_num:
        return False
    return {">": sv.lower() > c.lower(), "<": sv.lower() < c.lower(), ">=": sv.lower() >= c.lower(),
            "<=": sv.lower() <= c.lower()}[op]


def _flat(*args):
    out = []
    for a in args:
        if isinstance(a, Range):
            out.extend(a.flat())
        else:
            out.append(a)
    return out


def _as_range(v) -> Range:
    return v if isinstance(v, Range) else Range([[v]])


# --------------------------------------------------------------------------- evaluator
class WorkbookCalculator:
    """Evaluate cells of an openpyxl workbook loaded with ``data_only=False``."""

    def __init__(self, wb):
        self.wb = wb
        self.cache: dict[tuple[str, str], Any] = {}
        self.stack: set[tuple[str, str]] = set()

    # ---- cell access -------------------------------------------------------------------
    def _sheet(self, name: str | None, current: str):
        sname = current if name is None else name
        if sname not in self.wb.sheetnames:
            raise _Raise(REF)
        return self.wb[sname]

    def value(self, sheet: str, coord: str):
        key = (sheet, coord.replace("$", "").upper())
        if key in self.cache:
            return self.cache[key]
        if key in self.stack:
            raise _Raise(XLError("#CIRCULAR!"))
        ws = self.wb[sheet]
        raw = ws[key[1]].value
        if isinstance(raw, str) and raw.startswith("="):
            self.stack.add(key)
            try:
                v = self.evaluate_formula(raw, sheet)
            finally:
                self.stack.discard(key)
        else:
            v = raw
        self.cache[key] = v
        return v

    def evaluate_formula(self, formula: str, sheet: str):
        try:
            tree = parse(formula)
        except FormulaSyntaxError:
            return NAME
        try:
            v = self._ev(tree, sheet)
        except _Raise as r:
            return r.err
        if isinstance(v, Range):
            return v.rows[0][0] if v.rows and v.rows[0] else VALUE
        return v

    def _ref(self, text: str, sheet: str):
        sname = None
        if "!" in text:
            sname, text = text.rsplit("!", 1)
            if sname.startswith("'"):
                sname = sname[1:-1].replace("''", "'")
        ws = self._sheet(sname, sheet)
        target = ws.title
        text = text.replace("$", "").upper()
        if ":" in text:
            a, b = text.split(":")
            if a.isalpha() and b.isalpha():
                c1, c2 = column_index_from_string(a), column_index_from_string(b)
                r1, r2 = 1, ws.max_row
            else:
                m1 = re.match(r"([A-Z]+)(\d+)", a)
                m2 = re.match(r"([A-Z]+)(\d+)", b)
                c1, r1 = column_index_from_string(m1.group(1)), int(m1.group(2))
                c2, r2 = column_index_from_string(m2.group(1)), int(m2.group(2))
            rows = [[self.value(target, f"{get_column_letter(c)}{r}") for c in range(min(c1, c2), max(c1, c2) + 1)]
                    for r in range(min(r1, r2), max(r1, r2) + 1)]
            return Range(rows)
        return self.value(target, text)

    def _ev(self, n, sheet):
        kind = n[0]
        if kind == "num" or kind == "str" or kind == "bool":
            return n[1]
        if kind == "err":
            raise _Raise(n[1])
        if kind == "missing":
            return None
        if kind == "ref":
            return self._ref(n[1], sheet)
        if kind == "neg":
            return -_num(self._scalar(self._ev(n[1], sheet)))
        if kind == "pct":
            return _num(self._scalar(self._ev(n[1], sheet))) / 100
        if kind == "bin":
            op = n[1]
            a = self._scalar(self._ev(n[2], sheet))
            b = self._scalar(self._ev(n[3], sheet))
            if isinstance(a, XLError):
                raise _Raise(a)
            if isinstance(b, XLError):
                raise _Raise(b)
            if op == "&":
                return f"{_text(a)}{_text(b)}"
            if op in ("=", "<>", "<", ">", "<=", ">="):
                return _compare(op, a, b)
            x, y = _num(a), _num(b)
            if op == "+":
                return x + y
            if op == "-":
                return x - y
            if op == "*":
                return x * y
            if op == "/":
                if y == 0:
                    raise _Raise(DIV0)
                return x / y
            if op == "^":
                try:
                    return x ** y
                except (OverflowError, ZeroDivisionError):
                    raise _Raise(NUM)
        if kind == "func":
            return self._func(n[1], n[2], sheet)
        raise _Raise(VALUE)

    def _scalar(self, v):
        if isinstance(v, Range):
            if v.shape == (1, 1):
                return v.rows[0][0]
            raise _Raise(VALUE)
        return v

    # ---- functions -------------------------------------------------------------------
    def _func(self, name, args, sheet):
        if name not in SUPPORTED_FUNCTIONS:
            raise UnsupportedFormula(name)
        if name == "IFERROR":
            try:
                v = self._scalar(self._ev(args[0], sheet))
                if isinstance(v, XLError):
                    raise _Raise(v)
                return v
            except _Raise:
                return self._ev(args[1], sheet)
        if name == "IF":
            cond = self._scalar(self._ev(args[0], sheet))
            if isinstance(cond, XLError):
                raise _Raise(cond)
            truth = bool(_num(cond)) if not isinstance(cond, str) else bool(cond)
            if truth:
                return self._ev(args[1], sheet) if len(args) > 1 else True
            return self._ev(args[2], sheet) if len(args) > 2 else False
        ev = [self._ev(a, sheet) for a in args]
        if name == "SUM":
            # Ranges: numbers only (text/bool/blank ignored, errors propagate). Direct args: coerced.
            total = 0.0
            for e in ev:
                if isinstance(e, Range):
                    for v in e.flat():
                        if isinstance(v, XLError):
                            raise _Raise(v)
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            total += float(v)
                elif e is not None:
                    total += _num(e)
            return total
        if name in ("COUNT",):
            return float(sum(1 for v in _flat(*ev) if isinstance(v, (int, float, dt.date)) and not isinstance(v, bool)))
        if name == "COUNTA":
            return float(sum(1 for v in _flat(*ev) if v not in (None, "")))
        if name in ("AVERAGE",):
            vals = [_num(v) for v in _flat(*ev) if isinstance(v, (int, float)) and not isinstance(v, bool)]
            if not vals:
                raise _Raise(DIV0)
            return sum(vals) / len(vals)
        if name in ("MIN", "MAX"):
            vals = [_num(v) for v in _flat(*ev) if isinstance(v, (int, float)) and not isinstance(v, bool)]
            if not vals:
                return 0.0
            return min(vals) if name == "MIN" else max(vals)
        if name == "ABS":
            return abs(_num(self._scalar(ev[0])))
        if name in ("ROUND", "ROUNDUP", "ROUNDDOWN"):
            x = _num(self._scalar(ev[0]))
            d = int(_num(self._scalar(ev[1]))) if len(ev) > 1 else 0
            f = 10 ** d
            if name == "ROUND":
                return math.floor(abs(x) * f + 0.5) / f * (1 if x >= 0 else -1)
            if name == "ROUNDUP":
                return math.ceil(abs(x) * f - 1e-12) / f * (1 if x >= 0 else -1)
            return math.floor(abs(x) * f + 1e-12) / f * (1 if x >= 0 else -1)
        if name in ("AND", "OR"):
            vals = [bool(_num(v)) for v in _flat(*ev) if v is not None]
            return all(vals) if name == "AND" else any(vals)
        if name == "NOT":
            return not bool(_num(self._scalar(ev[0])))
        if name in ("SUMIF", "AVERAGEIF"):
            rng, crit = _as_range(ev[0]), self._scalar(ev[1])
            target = _as_range(ev[2]) if len(ev) > 2 and ev[2] is not None else rng
            pairs = [(t, c) for t, c in zip(target.flat(), rng.flat()) if _match_criterion(c, crit)]
            vals = [_num(t) for t, _c in pairs if isinstance(t, (int, float)) and not isinstance(t, bool)]
            if name == "SUMIF":
                return sum(vals)
            if not vals:
                raise _Raise(DIV0)
            return sum(vals) / len(vals)
        if name in ("SUMIFS", "AVERAGEIFS"):
            target = _as_range(ev[0]).flat()
            mask = self._mask(ev[1:], len(target))
            vals = [_num(t) for t, m in zip(target, mask) if m and isinstance(t, (int, float))
                    and not isinstance(t, bool)]
            if name == "SUMIFS":
                return sum(vals)
            if not vals:
                raise _Raise(DIV0)
            return sum(vals) / len(vals)
        if name == "COUNTIF":
            return float(sum(1 for v in _as_range(ev[0]).flat() if _match_criterion(v, self._scalar(ev[1]))))
        if name == "COUNTIFS":
            n = len(_as_range(ev[0]).flat())
            return float(sum(self._mask(ev, n)))
        if name == "SUMPRODUCT":
            arrays = [_as_range(e).flat() for e in ev]
            if len({len(a) for a in arrays}) != 1:
                raise _Raise(VALUE)
            total = 0.0
            for vals in zip(*arrays):
                p = 1.0
                for v in vals:
                    p *= _num(v) if isinstance(v, (int, float, bool)) else 0.0
                total += p
            return total
        if name == "XLOOKUP":
            key = self._scalar(ev[0])
            look, ret = _as_range(ev[1]).flat(), _as_range(ev[2])
            for i, v in enumerate(look):
                if _eq(v, key):
                    rr = ret.rows
                    if len(rr) == len(look):
                        return rr[i][0] if len(rr[i]) == 1 else Range([rr[i]])
                    return ret.flat()[i]
            if len(ev) > 3 and ev[3] is not None:
                return self._scalar(ev[3])
            raise _Raise(NA)
        if name == "VLOOKUP":
            key = self._scalar(ev[0])
            table = _as_range(ev[1]).rows
            col = int(_num(self._scalar(ev[2])))
            approx = bool(_num(self._scalar(ev[3]))) if len(ev) > 3 and ev[3] is not None else True
            if col < 1 or (table and col > len(table[0])):
                raise _Raise(REF)
            if not approx:
                for r in table:
                    if _eq(r[0], key):
                        return r[col - 1]
                raise _Raise(NA)
            best = None
            for r in table:
                if r[0] is not None and _compare("<=", r[0], key):
                    best = r
            if best is None:
                raise _Raise(NA)
            return best[col - 1]
        if name == "MATCH":
            key = self._scalar(ev[0])
            look = _as_range(ev[1]).flat()
            mtype = int(_num(self._scalar(ev[2]))) if len(ev) > 2 and ev[2] is not None else 1
            if mtype == 0:
                for i, v in enumerate(look):
                    if _eq(v, key):
                        return float(i + 1)
                raise _Raise(NA)
            best = None
            for i, v in enumerate(look):
                if v is None:
                    continue
                if (mtype == 1 and _compare("<=", v, key)) or (mtype == -1 and _compare(">=", v, key)):
                    best = i
            if best is None:
                raise _Raise(NA)
            return float(best + 1)
        if name == "INDEX":
            rng = _as_range(ev[0])
            r = int(_num(self._scalar(ev[1]))) if len(ev) > 1 and ev[1] is not None else 1
            c = int(_num(self._scalar(ev[2]))) if len(ev) > 2 and ev[2] is not None else 1
            rows, cols = rng.shape
            if rows == 1 and len(ev) == 2:
                r, c = 1, r
            if not (1 <= r <= rows and 1 <= c <= cols):
                raise _Raise(REF)
            return rng.rows[r - 1][c - 1]
        if name in ("MONTH", "YEAR", "DAY"):
            d = _to_date(self._scalar(ev[0]))
            return float({"MONTH": d.month, "YEAR": d.year, "DAY": d.day}[name])
        if name == "DATE":
            y, m, d = (int(_num(self._scalar(e))) for e in ev[:3])
            y += (m - 1) // 12
            m = (m - 1) % 12 + 1
            return float(_to_serial(dt.datetime(y, m, 1) + dt.timedelta(days=d - 1)))
        if name == "EOMONTH":
            d = _to_date(self._scalar(ev[0]))
            k = int(_num(self._scalar(ev[1])))
            mm = d.month - 1 + k
            y, m = d.year + mm // 12, mm % 12 + 1
            return float(_to_serial(dt.datetime(y, m, calendar.monthrange(y, m)[1])))
        if name == "NPV":
            from ..core.finance import npv
            r = _num(self._scalar(ev[0]))
            flows = [_num(v) for v in _flat(*ev[1:]) if isinstance(v, (int, float)) and not isinstance(v, bool)]
            return npv(r, [0.0] + flows)
        if name == "IRR":
            from ..core.finance import irr
            flows = [_num(v) for v in _flat(*ev) if isinstance(v, (int, float)) and not isinstance(v, bool)]
            res = irr(flows)
            if res.value is None:
                raise _Raise(NUM)
            return res.value
        if name == "PMT":
            from ..core.finance import pmt
            r, n, pv = (_num(self._scalar(e)) for e in ev[:3])
            return -pmt(r, int(n), pv)
        raise UnsupportedFormula(name)

    def _mask(self, pairs, n):
        mask = [True] * n
        for i in range(0, len(pairs) - 1, 2):
            rng = _as_range(pairs[i]).flat()
            crit = self._scalar(pairs[i + 1])
            if len(rng) != n:
                raise _Raise(VALUE)
            for j, v in enumerate(rng):
                if mask[j] and not _match_criterion(v, crit):
                    mask[j] = False
        return mask


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return _text(float(_to_serial(v)))
    return str(v)


def _eq(a, b) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        return str(a).lower() == str(b).lower() if a is not None and b is not None else False
    try:
        return abs(_cmp_key(a) - _cmp_key(b)) < 1e-12
    except TypeError:
        return False


def _compare(op, a, b) -> bool:
    if isinstance(a, str) and isinstance(b, str):
        a, b = a.lower(), b.lower()
    elif isinstance(a, str) or isinstance(b, str):
        if op in ("=", "<>"):
            return (op == "<>")
        # Excel orders text after numbers
        a_t, b_t = isinstance(a, str), isinstance(b, str)
        a, b = (1 if a_t else 0), (1 if b_t else 0)
    else:
        a = 0.0 if a is None else _cmp_key(a)
        b = 0.0 if b is None else _cmp_key(b)
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
