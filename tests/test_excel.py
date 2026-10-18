"""Excel integration: formula calculator, generated workbooks and the grader."""
import hashlib
import datetime as dt

import openpyxl
import pytest
from openpyxl import Workbook

from aurelia.core import finance as fin
from aurelia.excel import workbooks as W
from aurelia.excel.formula import WorkbookCalculator, functions_used, DIV0, UnsupportedFormula
from aurelia.excel.grader import grade_file, GradingError, record


def _calc(cells: dict, data=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for r in data or []:
        ws.append(r)
    for k, v in cells.items():
        ws[k] = v
    return WorkbookCalculator(wb)


class TestFormula:
    def test_arithmetic_and_precedence(self):
        c = _calc({"Z1": "=1+2*3", "Z2": "=(1+2)*3", "Z3": "=2^3^2", "Z4": "=-2^2", "Z5": "=10%*50",
                   "Z6": "=\"a\"&\"b\"&1"})
        assert c.value("S", "Z1") == 7
        assert c.value("S", "Z2") == 9
        assert c.value("S", "Z3") == 64          # Excel is left-associative
        assert c.value("S", "Z4") == 4           # unary minus binds tighter than ^ in Excel
        assert c.value("S", "Z5") == pytest.approx(5)
        assert c.value("S", "Z6") == "ab1"

    def test_errors_and_iferror(self):
        c = _calc({"Z1": "=1/0", "Z2": "=IFERROR(1/0,-1)", "Z3": "=IFERROR(Z1,\"x\")"})
        assert c.value("S", "Z1") == DIV0
        assert c.value("S", "Z2") == -1
        assert c.value("S", "Z3") == "x"

    def test_conditional_aggregates(self):
        data = [("seg", "rev", "d"), ("a", 10, dt.datetime(2027, 1, 3)), ("b", 20, dt.datetime(2027, 2, 3)),
                ("a", 30, dt.datetime(2027, 2, 9)), ("A", 5, None)]
        c = _calc({"Z1": "=SUMIFS(B2:B5,A2:A5,\"a\")", "Z2": "=COUNTIFS(B2:B5,\">=10\",A2:A5,\"a\")",
                   "Z3": "=SUMIFS(B:B,C:C,\">=\"&DATE(2027,2,1))", "Z4": "=AVERAGEIF(A2:A5,\"a\",B2:B5)",
                   "Z5": "=SUMIF(A2:A5,\"<>a\",B2:B5)", "Z6": "=SUMPRODUCT(B2:B3,B4:B5)"}, data)
        assert c.value("S", "Z1") == 45          # criteria are case-insensitive, as in Excel
        assert c.value("S", "Z2") == 2
        assert c.value("S", "Z3") == 50
        assert c.value("S", "Z4") == pytest.approx(15)
        assert c.value("S", "Z5") == 20
        assert c.value("S", "Z6") == 10 * 30 + 20 * 5

    def test_lookups(self):
        data = [("k", "v"), ("x", 1), ("y", 2), ("z", 3)]
        c = _calc({"Z1": "=XLOOKUP(\"y\",A2:A4,B2:B4)", "Z2": "=INDEX(B2:B4,MATCH(\"z\",A2:A4,0))",
                   "Z3": "=VLOOKUP(\"x\",A2:B4,2,FALSE)", "Z4": "=XLOOKUP(\"q\",A2:A4,B2:B4,\"none\")",
                   "Z5": "=IFERROR(MATCH(\"q\",A2:A4,0),0)"}, data)
        assert [c.value("S", f"Z{i}") for i in range(1, 6)] == [2, 3, 1, "none", 0]

    def test_finance_functions_match_core(self):
        c = _calc({"Z1": "=NPV(0.1,100,200)", "Z2": "=IRR(A1:A3)", "Z3": "=PMT(0.05,10,1000)"},
                  [(-100,), (60,), (60,)])
        assert c.value("S", "Z1") == pytest.approx(100 / 1.1 + 200 / 1.21)
        assert c.value("S", "Z2") == pytest.approx(fin.irr([-100, 60, 60]).value, abs=1e-9)
        assert c.value("S", "Z3") == pytest.approx(-fin.pmt(0.05, 10, 1000))
        assert c.value("S", "Z3") < 0            # Excel sign convention

    def test_dates(self):
        c = _calc({"A1": dt.datetime(2027, 2, 14), "Z1": "=MONTH(A1)", "Z2": "=YEAR(A1)",
                   "Z3": "=EOMONTH(A1,1)-EOMONTH(A1,0)", "Z4": "=DATE(2027,13,1)-DATE(2027,12,1)"})
        assert c.value("S", "Z1") == 2 and c.value("S", "Z2") == 2027
        assert c.value("S", "Z3") == 31
        assert c.value("S", "Z4") == 31

    def test_unsupported_and_circular(self):
        c = _calc({"Z1": "=OFFSET(A1,1,1)", "Z2": "=Z3+1", "Z3": "=Z2+1"})
        with pytest.raises(UnsupportedFormula):
            c.value("S", "Z1")
        assert str(c.value("S", "Z2")).startswith("#")

    def test_functions_used(self):
        assert functions_used("=IFERROR(_xlfn.XLOOKUP(A1,B:B,C:C),0)") == {"IFERROR", "XLOOKUP"}


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    d = tmp_path_factory.mktemp("wb")
    return d, W.generate_all(d)


class TestWorkbooks:
    def test_all_files_and_tabs(self, generated):
        d, paths = generated
        assert len(paths) == 8
        for p in paths:
            wb = openpyxl.load_workbook(p)
            assert "Instructions" in wb.sheetnames and "Assumptions" in wb.sheetnames
            assert W.identify(wb) is not None

    @pytest.mark.parametrize("key", ["revenue", "variance", "three_statement", "investment"])
    def test_reference_formulas_reproduce_python_answers(self, generated, key):
        d, _ = generated
        s = W.spec(key)
        wb = openpyxl.load_workbook(d / f"{s.file_stem}_REFERENCE.xlsx")
        calc = WorkbookCalculator(wb)
        for t in s.tasks:
            v = calc.value(t.sheet, t.cell)
            if isinstance(t.expected, str):
                assert v == t.expected, t.id
            else:
                assert v == pytest.approx(t.expected, abs=max(t.tolerance_abs, 1e-6)), t.id
        for r in s.reconciliations:
            ok, msg = r.check(calc.value)
            assert ok, msg

    def test_training_has_no_answers(self, generated):
        d, _ = generated
        for s in W.all_specs():
            wb = openpyxl.load_workbook(d / f"{s.file_stem}.xlsx")
            assert all(wb[t.sheet][t.cell].value is None for t in s.tasks)

    def test_hand_calculated_investment(self):
        s = W.spec("investment")
        npv = next(t for t in s.tasks if t.id == "4.1").expected
        flows = [-4_500_000, 1_050_000, 1_200_000, 1_300_000, 1_350_000, 1_400_000]
        assert npv == pytest.approx(sum(c / 1.1 ** i for i, c in enumerate(flows)))
        assert next(t for t in s.tasks if t.id == "4.4").expected == 4   # cumulative turns positive in year 4

    def test_never_overwrites(self, tmp_path):
        p1 = W.generate("investment", tmp_path)
        p2 = W.generate("investment", tmp_path)
        assert p1 != p2 and p1.exists() and p2.exists()
        assert "(2)" in p2.name

    def test_three_statement_balances_by_hand(self):
        e = W.three_statement_expected()
        assert all(abs(x) < 1e-6 for x in e["check"])
        # month 1 by hand: EBITDA 6.8m × 26% = 1,768,000; EBIT 918,000; tax 183,600
        assert e["ebitda"][0] == pytest.approx(1_768_000)
        assert e["tax"][0] == pytest.approx(183_600)


def _sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


class TestGrader:
    def test_reference_scores_full_marks_and_original_untouched(self, generated, tmp_path):
        d, _ = generated
        for s in W.all_specs():
            src = d / f"{s.file_stem}_REFERENCE.xlsx"
            before = _sha(src)
            rep = grade_file(src, tmp_path / "graded")
            assert rep.score == pytest.approx(100.0), (s.key, [t for t in rep.tasks if t.status != "correct"][:3])
            assert all(ok for _, ok, _ in rep.reconciliations)
            assert _sha(src) == before
            assert rep.graded_copy != str(src)

    def test_blank_scores_zero(self, generated, tmp_path):
        d, _ = generated
        rep = grade_file(d / "01_Revenue_KPIs.xlsx", tmp_path)
        assert rep.score == 0
        assert all(t.status == "missing" for t in rep.tasks)

    def test_typed_values_get_value_credit_only(self, generated, tmp_path):
        d, _ = generated
        s = W.spec("investment")
        wb = openpyxl.load_workbook(d / f"{s.file_stem}.xlsx")
        for t in s.tasks:
            wb[t.sheet][t.cell] = t.expected
        p = tmp_path / "typed.xlsx"
        wb.save(p)
        rep = grade_file(p, tmp_path / "g")
        assert rep.score == pytest.approx(70.0)
        assert any("typed value" in f for t in rep.tasks for f in t.feedback)

    def test_common_mistake_feedback(self, generated, tmp_path):
        d, _ = generated
        wb = openpyxl.load_workbook(d / "04_Investment_Appraisal.xlsx")
        wb["Exercises"]["D5"] = "=NPV(Assumptions!B4,Cashflows!C5:H5)"     # year 0 inside NPV()
        wb["Exercises"]["D11"] = "=PMT(Assumptions!B6,Assumptions!B7,Assumptions!B5)"   # sign not flipped
        p = tmp_path / "mistakes.xlsx"
        wb.save(p)
        rep = grade_file(p, tmp_path / "g")
        t = {x.id: x for x in rep.tasks}
        assert t["4.1"].status == "wrong" and "year 0 inside NPV" in t["4.1"].feedback[0]
        assert "flip the sign" in t["4.7"].feedback[0]

    def test_wrong_denominator_feedback(self, generated, tmp_path):
        d, _ = generated
        wb = openpyxl.load_workbook(d / "01_Revenue_KPIs.xlsx")
        ex = wb["Exercises"]
        ex["D5"] = "=SUM(Bookings!F:F)"
        ex["D6"] = "=SUM(Bookings!D:D)"
        ex["D7"] = "=SUM(Daily!B:B)"
        ex["D9"] = "=D5/D7"       # RevPAR typed as ADR
        p = tmp_path / "adr.xlsx"
        wb.save(p)
        rep = grade_file(p, tmp_path / "g")
        t = {x.id: x for x in rep.tasks}
        assert t["1.1"].status == "correct" and t["1.3"].status == "correct"
        assert "RevPAR" in t["1.5"].feedback[0]

    def test_rejects_foreign_files(self, tmp_path):
        p = tmp_path / "other.xlsx"
        Workbook().save(p)
        with pytest.raises(GradingError):
            grade_file(p, tmp_path / "g")
        with pytest.raises(GradingError):
            grade_file(tmp_path / "missing.xlsx", tmp_path / "g")
        bad = tmp_path / "bad.xlsx"
        bad.write_bytes(b"not a zip")
        with pytest.raises(GradingError):
            grade_file(bad, tmp_path / "g")

    def test_record(self, generated, tmp_path):
        from aurelia.persistence.db import Database
        d, _ = generated
        db = Database(tmp_path / "t.db")
        rep = grade_file(d / "04_Investment_Appraisal_REFERENCE.xlsx", tmp_path / "g")
        record(db, rep)
        assert db.one("SELECT score FROM workbook_grades")["score"] == pytest.approx(100.0)
        db.close()
