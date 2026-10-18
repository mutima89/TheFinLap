"""Decision Room, learning path, coach, competency rules, capstone, reports, portfolio and persistence."""
import json
import sqlite3

import pytest

from aurelia.analysis import decisions as D
from aurelia.analysis import capstone as K
from aurelia.learning import content as C
from aurelia.learning import coach
from aurelia.learning import progress as P
from aurelia.learning.glossary import search
from aurelia.persistence.db import Database


@pytest.fixture
def db(tmp_path):
    d = Database(tmp_path / "t.db")
    yield d
    d.close()


# --------------------------------------------------------------------------- Decision Room
def test_nine_cases_with_complete_structure():
    assert len(D.CASES) == 9
    for c in D.CASES:
        assert c.question and c.situation and len(c.options) >= 2
        assert any(r for _d, r in c.drivers) and any(not r for _d, r in c.drivers)
        assert c.investigate_next and c.risk_terms
        assert max(o.quality for o in c.options) >= 0.9


@pytest.mark.parametrize("case", D.CASES, ids=lambda c: c.id)
def test_option_components_sum_to_total(case):
    for o in case.options:
        r = D.evaluate_option(case, o)
        assert sum(v for _n, v in r.components) == pytest.approx(r.total, rel=1e-9, abs=1.0)


def _strong(case):
    best = max(case.options, key=lambda o: o.quality)
    truth = D.evaluate_option(case, best).total
    text = (f"I recommend {best.label.lower()} because the analysis shows {truth:,.0f} for the decision metric. "
            f"The main risk is {case.risk_terms[0]}; the key assumption could be wrong, so we will monitor it monthly. "
            "Next step: the finance manager owns the follow-up and reports by the end of the month with the "
            "updated numbers and a decision on whether to continue.")
    return best, D.Submission(case.id, best.id, [d for d, r in case.drivers if r], truth, text)


@pytest.mark.parametrize("case", D.CASES, ids=lambda c: c.id)
def test_strong_submission_scores_high_and_weak_scores_low(case):
    best, sub = _strong(case)
    g = D.grade(case, sub)
    assert g.overall >= 80, g.feedback
    worst = min(case.options, key=lambda o: o.quality)
    weak = D.Submission(case.id, worst.id, [d for d, r in case.drivers if not r], None, "ok fine")
    gw = D.grade(case, weak)
    assert gw.overall < 40
    assert set(g.scores) == set(D.WEIGHTS)
    assert sum(D.WEIGHTS.values()) == pytest.approx(1.0)


def test_wrong_sign_estimate_gets_no_numerical_credit():
    case = D.CASE_BY_ID["ota_campaign"]
    best, sub = _strong(case)
    sub.estimate = -sub.estimate
    assert D.grade(case, sub).scores["numerical"] == 0.0


def test_condition_terms_lift_a_defensible_option():
    case = D.CASE_BY_ID["renovation"]
    opt = next(o for o in case.options if o.condition_terms and o.quality < 0.9)
    base = D.Submission(case.id, opt.id, [], 0, "I recommend this option.")
    with_cond = D.Submission(case.id, opt.id, [], 0, "I recommend this option " + opt.condition_terms[0] + ".")
    assert D.grade(case, with_cond).scores["judgement"] > D.grade(case, base).scores["judgement"]


@pytest.mark.parametrize("case", D.CASES, ids=lambda c: c.id)
def test_outcome_is_reported_with_next_investigations(case):
    o = D.outcome(case, case.options[0])
    assert {"expected", "realised", "investigate_next"} <= set(o)
    assert o["investigate_next"]


# --------------------------------------------------------------------------- learning content
def test_nine_stages_and_every_topic_complete():
    assert len(C.STAGES) == 9
    by_stage = C.topics_by_stage()
    assert set(by_stage) == set(range(1, 10))
    for t in C.topics():
        for f in ("explanation", "context", "mistake", "interpretation"):
            assert getattr(t, f).strip(), (t.id, f)
        assert len(t.worked()) > 120, t.id
        assert t.guided and t.independent and t.applied and t.check
        assert t.competency in P.COMPETENCIES


def test_every_exercise_accepts_its_own_answer_and_known_mistakes_are_diagnosed():
    for t in C.topics():
        exs = [t.guided, t.independent, *t.check] + ([t.applied.exercise] if t.applied.exercise else [])
        for ex in exs:
            ans = ex.answer if ex.unit == "choice" else float(ex.answer)
            assert C.is_correct(ex, ans), ex.id
            if ex.unit == "choice":
                assert ex.answer in ex.choices, ex.id
            for val, tag, _msg in ex.mistakes:
                d = coach.diagnose(ex, val)
                assert not d.correct and d.tag == tag, (ex.id, val, d)


def test_parse_answer_formats():
    assert C.parse_answer("12.5%", "pct") == pytest.approx(0.125)
    assert C.parse_answer("0.125", "pct") == pytest.approx(0.125)
    assert C.parse_answer("SAR 1,234.5", "currency") == pytest.approx(1234.5)
    assert C.parse_answer("abc", "number") is None


def test_coach_diagnoses_generic_errors_and_hints_progress():
    ex = C.Exercise("x", "q", 250_000.0, "currency", 0.001, hints=["first", "second"], solution=["do it"])
    assert coach.diagnose(ex, "-250000").tag == "sign_error"
    assert coach.diagnose(ex, "250").tag == "units_scale"
    assert coach.diagnose(ex, "hello").tag == "unparseable"
    p = C.Exercise("p", "q", 0.12, "pct", 0.001)
    assert coach.diagnose(p, 0.0012).tag == "percent_scale"
    assert coach.hint(ex, 1) == "first" and coach.hint(ex, 2) == "second" and "do it" in coach.hint(ex, 3)


# --------------------------------------------------------------------------- competency rules
def test_quiz_evidence_alone_is_capped_at_developing(db):
    for i in range(10):
        P.record_attempt(db, "check", "s1_pnl", f"q{i}", 1.0, "fpa")
    fpa = next(c for c in P.competency_profile(db) if c.key == "fpa")
    assert fpa.level == "Developing" and fpa.score == pytest.approx(1.0)


def test_demonstrated_needs_two_practical_passes_one_hint_free(db):
    P.record_attempt(db, "independent", "s4_variance", "a", 0.9, "fpa", hints_used=2)
    P.record_attempt(db, "applied", "s4_variance", "b", 0.9, "fpa", hints_used=1)
    assert next(c for c in P.competency_profile(db) if c.key == "fpa").level == "Developing"
    P.record_attempt(db, "decision", "fpa_gap", "c", 0.8, "fpa", hints_used=0)
    assert next(c for c in P.competency_profile(db) if c.key == "fpa").level == "Demonstrated"


def test_recurring_mistakes_are_reported(db):
    for _ in range(3):
        P.record_attempt(db, "guided", "s4_variance", "a", 0.0, "fpa", mistake_tag="sign_error")
    fpa = next(c for c in P.competency_profile(db) if c.key == "fpa")
    assert ("sign_error", 3) in fpa.recurring_mistakes
    assert coach.gaps(db)


def test_steps_unlock_in_order_and_stages_need_applied_and_check(db):
    tid = "s1_pnl"
    assert P.step_unlocked(db, tid, "guided") and not P.step_unlocked(db, tid, "independent")
    P.mark_step(db, tid, "guided", 0.5)
    assert not P.step_unlocked(db, tid, "independent")
    P.mark_step(db, tid, "guided", 0.9)
    assert P.step_unlocked(db, tid, "independent")
    tbs = C.topics_by_stage()
    for t in tbs[1]:
        P.mark_step(db, t, "check", 1.0)
    assert not P.stage_unlocked(db, 2, tbs), "a check without an applied task must not unlock"
    for t in tbs[1]:
        P.mark_step(db, t, "applied", 1.0)
    assert P.stage_unlocked(db, 2, tbs) and not P.stage_unlocked(db, 3, tbs)


# --------------------------------------------------------------------------- capstone
def test_capstone_correct_rules_reproduce_controls_and_wrong_rules_do_not():
    raw = K.messy_rooms()
    ctrl = K.hotel_controls()
    clean = K.apply_rules(raw, K.correct_rules())
    assert clean["room_nights"].sum() == pytest.approx(ctrl["occupied_rooms"].sum(), abs=0.5)
    assert clean["room_revenue"].sum() == pytest.approx(ctrl["room_revenue"].sum(), abs=1.0)
    bad = K.apply_rules(raw, {"drop_any"})
    assert abs(bad["room_revenue"].sum() - ctrl["room_revenue"].sum()) > 1000
    assert K.grade_step(1, {"rules": list(K.correct_rules()), "duplicates": K.DUPLICATES}).score == pytest.approx(1)
    assert K.grade_step(1, {"rules": ["drop_any"], "duplicates": 0}).score < 0.5


def test_capstone_expected_answers_score_full_marks_and_overall_needs_all_steps(db):
    e = K.expected()
    k = e["kpis"]
    answers = {
        1: {"rules": list(K.correct_rules()), "duplicates": K.DUPLICATES},
        2: {"occupancy": k["occupancy"], "adr": k["adr"], "revpar": k["revpar"], "net_adr": k["net_adr"]},
        3: {"rev_var": e["rev_var"], "rate_var": e["rate_var"], "vol_var": e["vol_var"],
            "fu": "F" if e["rev_var"] >= 0 else "U"},
        4: {"method": "Seasonal (same quarter last year × growth)", "forecast": e["forecast_truth"]},
        5: {"ebitda": e["ebitda"], "cfo": e["cfo"], "gap_item": e["gap_item"]},
        6: {"npv_a": e["npv_a"], "npv_b": e["npv_b"], "choice": "A" if e["npv_a"] > e["npv_b"] else "B", "why": ""},
        7: {"down_delta": e["down_ebitda_delta"], "down_pct": e["down_pct"],
            "insight": "Fixed costs do not fall with demand: operating leverage."},
        8: {"kpis": ["EBITDA vs budget", "Closing cash", "RevPAR", "GOP margin", "Net RevPAR", "Labour cost %"]},
    }
    for step, ans in answers.items():
        r = K.grade_step(step, ans)
        assert r.score >= 0.95, (step, r.notes)
        K.save_step(db, step, ans, r)
    assert K.overall_score(db) is None
    for step, ans in ((9, {"text": "x"}), (10, {"answers": [1, 1, 1, 1], "text": ""})):
        K.save_step(db, step, ans, K.grade_step(step, ans))
    assert K.overall_score(db) is not None
    assert K.grade_step(8, {"kpis": ["Website page views", "Social media followers"]}).score < 0.5


def test_capstone_dataset_export_never_overwrites(tmp_path):
    a = K.export_dataset(tmp_path)
    b = K.export_dataset(tmp_path)
    assert a and set(a).isdisjoint(b)


# --------------------------------------------------------------------------- reports and portfolio
def test_management_report_has_seven_sections_and_marks_unconfirmed_causes():
    from aurelia.analysis.fpa import build_fpa, actual_world_assumptions
    from aurelia.model.assumptions import baseline
    from aurelia.reports.management_report import build_report
    h = build_report(build_fpa(baseline(), actual_world_assumptions()))
    for i, sec in enumerate(["Executive summary", "Actual, budget and forecast", "Material variances",
                             "Evidence-supported explanations", "Forward-looking risks",
                             "Recommended corrective actions", "Action owners and due dates"], 1):
        assert f"{i}. {sec}" in h
    assert "confirm" in h.lower()
    assert "fictional" in h.lower()


def test_portfolio_export_with_disclaimer(db, tmp_path, qapp):
    from aurelia.reports import portfolio as PF
    case = D.CASE_BY_ID["ota_campaign"]
    _best, sub = _strong(case)
    g = D.grade(case, sub)
    db.execute("INSERT INTO decisions (ts, case_id, option_id, justification, scores) VALUES (?,?,?,?,?)",
               ("2026-01-01T00:00:00", case.id, sub.option_id, sub.justification,
                json.dumps({"overall": g.overall, "dimensions": g.scores})))
    P.record_decision(db, case, g)
    paths = PF.export(db, tmp_path, learner="Test learner")
    assert {p.suffix for p in paths} == {".html", ".json", ".pdf"}
    html = paths[0].read_text(encoding="utf-8")
    assert "not an accredited qualification" in html and "Test learner" in html and case.title in html
    data = json.loads(paths[1].read_text(encoding="utf-8"))
    assert data["decisions"][0]["case"] == case.title


def test_statement_export_round_trip(tmp_path, base_result):
    from openpyxl import load_workbook
    from aurelia.reports.excel_export import export_statements
    p = export_statements(base_result, tmp_path / "s.xlsx", "Base case")
    wb = load_workbook(p)
    assert {"Income statement", "Balance sheet", "Cash flow", "Integrity checks", "About"} <= set(wb.sheetnames)
    assert "Fictional" in wb["About"]["A1"].value


# --------------------------------------------------------------------------- persistence
def test_backup_validate_restore_and_reset(db, tmp_path):
    from aurelia.model.scenarios import ScenarioStore
    store = ScenarioStore(db)
    from aurelia.model.assumptions import baseline
    store.create("Mine", baseline(), "kept")
    bk = db.backup(tmp_path / "b.db")
    assert db.validate_backup(bk) == []
    safety = db.reset_all()
    assert safety.exists()
    assert "Mine" not in ScenarioStore(db).names()
    db.restore(bk)
    assert "Mine" in ScenarioStore(db).names()


def test_invalid_backup_is_rejected(db, tmp_path):
    junk = tmp_path / "junk.db"
    junk.write_bytes(b"not a database")
    assert db.validate_backup(junk)
    other = tmp_path / "other.db"
    sqlite3.connect(other).execute("CREATE TABLE x (a)").connection.commit()
    assert db.validate_backup(other)


def test_integrity_check_clean_on_new_database(db):
    assert db.integrity_check() == []


# --------------------------------------------------------------------------- search
def test_search_finds_across_sources():
    kinds = {h.kind for h in search("revpar")}
    assert {"Glossary", "KPI", "Lesson"} <= kinds
    assert any(h.kind == "Decision case" for h in search("refurbishment"))
    assert search("") == []
